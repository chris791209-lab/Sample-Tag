"""Card layout shared by the workbook builder and the PNG renderer: which
template cell carries which field, plus sheet geometry read from the raw XML
(openpyxl mis-reads grouped <col min=.. max=..> widths, so we parse them)."""
from __future__ import annotations

import re
from datetime import date

from openpyxl.utils import column_index_from_string, get_column_letter

from .config import BLOCK_ROWS, BLOCK_START, CARD_COLS, LEFT_COL, RIGHT_COL

# (row offset within block, col offset within card) -> field.
# Row offset 0 = template row 2, col offset 0 = column B (left) / H (right).
FIELD_CELLS = {
    (1, 0): "program",
    (2, 1): "description",
    (2, 4): "sample_date",
    (5, 1): "style",
    (5, 4): "pcn_item_id",
    (6, 1): "apex_id",
    (6, 4): "spark_pid",
    (7, 1): "first_cost",
    (7, 4): "country",
    (8, 1): "retail",
    (8, 4): "duty",
    (12, 1): "ly_dpci",
    (13, 1): "ly_cost",
    (14, 1): "ly_retail",
    (15, 1): "ly_duty",
    (19, 1): "comments",
}
NUMERIC_FIELDS = {"first_cost", "retail", "duty", "ly_cost", "ly_retail", "ly_duty"}
TEXT_FIELDS = {"program", "description", "style", "pcn_item_id", "apex_id", "spark_pid",
               "country", "ly_dpci", "comments"}

VENDOR_CELL = (0, 0)          # B2 = Instructions!$B$4
THUMB_CELL = (19, 0)          # B21 placeholder "[INSERT THUMBNAIL IMAGE HERE]"
BARCODE_BOX = ((12, 2), (17, 4))   # D14:F19 (rows 12..17, cols 2..4)
THUMB_BOX = ((19, 0), (23, 0))     # B21:B25
IMU_CELLS = {(9, 1): ("imu", False), (10, 1): ("imu", True),
             (16, 1): ("ly_imu", False), (17, 1): ("ly_imu", True)}
RETAIL_ECHO = (19, 4)         # F21 = C10

EXCEL_EPOCH = date(1899, 12, 30)
EMU_PER_PX = 9525
EMU_PER_PT = 12700


def excel_serial(d) -> int | None:
    if d is None or d == "":
        return None
    if hasattr(d, "date") and not isinstance(d, date):
        d = d.date()
    try:
        return (d - EXCEL_EPOCH).days
    except TypeError:
        return None


def card_origin(index: int) -> tuple[int, int]:
    """card #index (0-based) -> (top row, left column) in the generated sheet."""
    block, side = divmod(index, 2)
    return BLOCK_START + block * BLOCK_ROWS, (LEFT_COL if side == 0 else RIGHT_COL)


class Geometry:
    """Column widths (chars) and row heights (points) of the template sheet."""

    def __init__(self, sheet_xml: str):
        fmt = re.search(r"<sheetFormatPr\b[^>]*>", sheet_xml)
        fmt = fmt.group(0) if fmt else ""
        dcw = re.search(r'defaultColWidth="([0-9.]+)"', fmt)
        drh = re.search(r'defaultRowHeight="([0-9.]+)"', fmt)
        self.default_col = float(dcw.group(1)) if dcw else 8.43
        self.default_row = float(drh.group(1)) if drh else 15.0
        self.cols: dict[int, float] = {}
        for m in re.finditer(r"<col\b[^>]*/>", sheet_xml):
            t = m.group(0)
            lo = int(re.search(r'min="(\d+)"', t).group(1))
            hi = int(re.search(r'max="(\d+)"', t).group(1))
            w = re.search(r'width="([0-9.]+)"', t)
            hidden = 'hidden="1"' in t
            for c in range(lo, min(hi, 60) + 1):
                self.cols[c] = 0.0 if hidden else (float(w.group(1)) if w else self.default_col)
        self.rows: dict[int, float] = {}
        for m in re.finditer(r"<row\b[^>]*>", sheet_xml):
            t = m.group(0)
            r = int(re.search(r'\br="(\d+)"', t).group(1))
            h = re.search(r'\bht="([0-9.]+)"', t)
            if h:
                self.rows[r] = float(h.group(1))

    def col_px(self, c: int) -> int:
        return int(round(self.cols.get(c, self.default_col) * 7))

    def row_pt(self, r: int) -> float:
        return self.rows.get(r, self.default_row)

    def row_px(self, r: int) -> float:
        return self.row_pt(r) * 96 / 72


def ref(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def split_ref(a1: str) -> tuple[int, int]:
    m = re.fullmatch(r"\$?([A-Z]{1,3})\$?(\d+)", a1)
    return int(m.group(2)), column_index_from_string(m.group(1))


def parse_range(rng: str) -> tuple[int, int, int, int]:
    """'B2:F2' -> (r1, c1, r2, c2)"""
    if ":" in rng:
        a, b = rng.split(":")
    else:
        a = b = rng
    r1, c1 = split_ref(a)
    r2, c2 = split_ref(b)
    return r1, c1, r2, c2


def in_proto_card(rng: str) -> bool:
    r1, c1, r2, c2 = parse_range(rng)
    return (BLOCK_START <= r1 and r2 <= BLOCK_START + BLOCK_ROWS - 1
            and LEFT_COL <= c1 and c2 <= LEFT_COL + CARD_COLS - 1)


def shift_range(rng: str, dr: int, dc: int) -> str:
    r1, c1, r2, c2 = parse_range(rng)
    a, b = ref(r1 + dr, c1 + dc), ref(r2 + dr, c2 + dc)
    return a if a == b else f"{a}:{b}"


_REF = re.compile(r"(?<![A-Za-z0-9_!$'\].])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![0-9A-Za-z_(!])")


def shift_formula(f: str, dr: int, dc: int) -> str:
    """Shift relative A1 references (respecting $) outside string literals and
    leaving sheet-qualified references (Instructions!$B$4) alone."""
    if not f or (dr == 0 and dc == 0):
        return f
    parts = re.split(r'("(?:[^"]|"")*")', f)
    out = []
    for p in parts:
        if p.startswith('"'):
            out.append(p)
            continue

        def sub(m):
            cabs, col, rabs, row = m.groups()
            c = column_index_from_string(col)
            r = int(row)
            if not cabs:
                c += dc
            if not rabs:
                r += dr
            return f"{cabs}{get_column_letter(c)}{rabs}{r}"
        out.append(_REF.sub(sub, p))
    return "".join(out)
